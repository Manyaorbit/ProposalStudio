"""Turns a text-based PDF sample into an editable Word document (in memory).

This is a best-effort rebuild: headings, paragraphs, bullets, fonts, sizes,
colours and page size are recovered from the PDF. Images, tables, multi-column
layouts and exact positioning are not. A .docx sample always gives the best result.
"""
import io
import re
from collections import Counter

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from docengine import DocError

MAX_PAGES = 60
BULLET_RE = re.compile(r"^[•●▪◦‣⁃∙·\-–—*]\s+")
BULLET_STRIP = re.compile(r"^[•●▪◦‣⁃∙·\-–—*]\s*")
END_PUNCT = (".", ",", ";", ":", "!", "?")


def _clean_font(name: str) -> str:
    base = re.sub(r"^[A-Z]{6}\+", "", name or "")
    base = re.split(r"[-,]", base)[0]
    base = re.sub(r"(PSMT|MT|PS)$", "", base)
    base = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base).strip()
    return base or "Calibri"


def _is_bold(name):
    n = (name or "").lower()
    return any(k in n for k in ("bold", "black", "heavy", "semibold", "demi"))


def _is_italic(name):
    n = (name or "").lower()
    return "italic" in n or "oblique" in n


def _color_hex(c):
    try:
        if c is None:
            return None
        if isinstance(c, (int, float)):
            c = (c,)
        c = tuple(c)
        if len(c) == 1:
            r = g = b = c[0]
        elif len(c) == 3:
            r, g, b = c
        elif len(c) == 4:
            cc, m, y, k = c
            r, g, b = (1 - cc) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)
        else:
            return None
        return "%02X%02X%02X" % tuple(max(0, min(255, round(v * 255))) for v in (r, g, b))
    except Exception:
        return None


class Line:
    def __init__(self, words, page_no, page_w, page_h):
        self.words = sorted(words, key=lambda w: w["x0"])
        self.page_no, self.page_w, self.page_h = page_no, page_w, page_h
        self.text = " ".join(w["text"] for w in self.words)
        self.top = min(w["top"] for w in self.words)
        self.bottom = max(w["bottom"] for w in self.words)
        self.x0 = min(w["x0"] for w in self.words)
        self.x1 = max(w["x1"] for w in self.words)
        weight = Counter()
        fonts, colors = Counter(), Counter()
        bold_w = 0
        for w in self.words:
            n = len(w["text"])
            weight[round(w["size"] * 2) / 2] += n
            fonts[_clean_font(w["fontname"])] += n
            colors[_color_hex(w.get("non_stroking_color"))] += n
            if _is_bold(w["fontname"]):
                bold_w += n
        total = max(1, sum(len(w["text"]) for w in self.words))
        self.size = weight.most_common(1)[0][0]
        self.font = fonts.most_common(1)[0][0]
        self.color = colors.most_common(1)[0][0]
        self.bold = bold_w / total >= 0.7
        self.centered = abs((self.x0 + self.x1) / 2 - page_w / 2) < 14 and self.x0 > 60

    def runs(self):
        out = []
        for w in self.words:
            b, i = _is_bold(w["fontname"]), _is_italic(w["fontname"])
            if out and out[-1][1] == b and out[-1][2] == i:
                out[-1][0] += " " + w["text"]
            else:
                out.append([(" " if out else "") + w["text"], b, i])
        return out


def _page_lines(page, page_no):
    words = page.extract_words(extra_attrs=["fontname", "size", "non_stroking_color"],
                               keep_blank_chars=False, x_tolerance=1.5)
    words = [w for w in words if w["text"].strip()]
    words.sort(key=lambda w: (w["top"], w["x0"]))
    lines, cur, cur_top = [], [], None
    for w in words:
        if cur and abs(w["top"] - cur_top) <= 0.4 * w["size"]:
            cur.append(w)
        else:
            if cur:
                lines.append(Line(cur, page_no, float(page.width), float(page.height)))
            cur, cur_top = [w], w["top"]
    if cur:
        lines.append(Line(cur, page_no, float(page.width), float(page.height)))
    return lines


def _drop_headers_footers(lines, n_pages):
    def norm(t):
        return re.sub(r"\d+", "#", t.lower()).strip()

    def in_margin(l):
        return l.top < 0.08 * l.page_h or l.bottom > 0.92 * l.page_h

    keep = []
    counts = Counter(norm(l.text) for l in lines if in_margin(l))
    for l in lines:
        if in_margin(l):
            if re.fullmatch(r"(page\s*)?\d+(\s*(of|/)\s*\d+)?", l.text.strip(), re.I):
                continue
            if n_pages >= 3 and counts[norm(l.text)] >= max(2, n_pages // 2):
                continue
        keep.append(l)
    return keep


def _set_style_font(style, name, size, bold=None, color=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is not None:
        for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            k = qn("w:" + a)
            if k in rf.attrib:
                del rf.attrib[k]


def pdf_to_docx(data: bytes):
    """Return (docx_bytes, notes)."""
    if not data.startswith(b"%PDF-"):
        raise DocError("That doesn't look like a PDF file.")
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception:
        raise DocError("That PDF couldn't be opened (it may be damaged or password-protected).")
    with pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise DocError(f"PDFs longer than {MAX_PAGES} pages aren't supported.")
        all_lines, page_w, page_h = [], 595.0, 842.0
        for i, page in enumerate(pdf.pages):
            if i == 0:
                page_w, page_h = float(page.width), float(page.height)
            all_lines.extend(_page_lines(page, i))
        n_pages = len(pdf.pages)

    if not all_lines:
        raise DocError("No selectable text found in that PDF (scanned PDFs aren't supported). "
                       "Please upload the Word version instead.")
    lines = _drop_headers_footers(all_lines, n_pages)
    if not lines:
        raise DocError("No usable text found in that PDF.")

    # Body style = most common size/font/colour by amount of text
    size_w, font_w, col_w = Counter(), Counter(), Counter()
    for l in lines:
        n = len(l.text)
        size_w[l.size] += n
        font_w[l.font] += n
        col_w[l.color] += n
    body_size = size_w.most_common(1)[0][0]
    body_font = font_w.most_common(1)[0][0]
    body_color = col_w.most_common(1)[0][0] or "000000"

    big = sorted({l.size for l in lines if l.size >= body_size + 1.0}, reverse=True)
    title_size = None
    if big:
        top_lines = [l for l in lines if l.size == big[0]]
        if len(top_lines) <= 3 and all(l.page_no == 0 for l in top_lines):
            title_size = big[0]
    # A size only counts as a heading level if it repeats (real headings recur; a one-off larger line
    # such as a cover-page subtitle does not). One-offs stay ordinary text with direct formatting.
    size_counts = Counter(l.size for l in lines if len(l.text) <= 90 and l.size in big and l.size != title_size)
    h_tiers = [s for s in big if s != title_size and size_counts[s] >= 2]

    def kind_of(l):
        if title_size is not None and l.size == title_size:
            return "title"
        if l.size in h_tiers:
            return "h1" if h_tiers.index(l.size) == 0 else "h2"
        if BULLET_RE.match(l.text):
            return "bullet"
        if (l.bold and len(l.text) <= 70 and not l.text.endswith(END_PUNCT)):
            return "h2" if h_tiers else "h1"
        return "body"

    # Typical distance between two wrapped lines of the same paragraph (top to top)
    pitches = sorted(b.top - a.top for a, b in zip(lines, lines[1:])
                     if a.page_no == b.page_no and kind_of(a) == "body" and kind_of(b) == "body"
                     and 0 < b.top - a.top < 2.0 * body_size)
    line_pitch = pitches[len(pitches) // 2] if pitches else 1.2 * body_size
    para_break = line_pitch + 0.35 * body_size

    # Group lines into paragraphs
    paras = []  # dict(kind, lines)
    prev = None
    for l in lines:
        k = kind_of(l)
        new = True
        if prev is not None and paras:
            last = paras[-1]
            gap = l.top - prev.bottom
            same_page = l.page_no == prev.page_no
            if k in ("title", "h1", "h2") and last["kind"] == k and l.size == prev.size and gap < 0.8 * l.size:
                new = False
            elif k == "body" and last["kind"] == "body":
                indented_start = prev.text.endswith(END_PUNCT) and l.x0 > prev.x0 + 6
                if same_page:
                    new = (l.top - prev.top) > para_break or indented_start
                else:
                    new = prev.text.endswith((".", "!", "?", ":")) or indented_start
        if new:
            paras.append({"kind": k, "lines": [l]})
        else:
            paras[-1]["lines"].append(l)
        prev = l

    # Build the Word file
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(page_w), Pt(page_h)
    body_lines = [l for l in lines if kind_of(l) in ("body", "bullet")] or lines
    left = max(36, min(l.x0 for l in body_lines))
    right = max(36, page_w - max(l.x1 for l in body_lines))
    top = max(36, min(l.top for l in lines))
    bottom = max(36, page_h - max(l.bottom for l in lines))
    sec.left_margin, sec.right_margin = Pt(min(left, 130)), Pt(min(right, 130))
    sec.top_margin, sec.bottom_margin = Pt(min(top, 110)), Pt(min(bottom, 110))

    styles = doc.styles
    normal = styles["Normal"]
    _set_style_font(normal, body_font, body_size, color=body_color)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    def head_look(kind):
        sample = next((l for l in lines if kind_of(l) == kind), None)
        if sample is None:
            return None
        return sample

    for kind, style_name in (("title", "Title"), ("h1", "Heading 1"), ("h2", "Heading 2")):
        l = head_look(kind)
        if l is None:
            continue
        st = styles[style_name]
        col = l.color if l.color and l.color != "FFFFFF" else "000000"
        _set_style_font(st, l.font, l.size, bold=l.bold, color=col)
        ppr = st.element.pPr
        if ppr is not None:
            for bdr in ppr.findall(qn("w:pBdr")):
                ppr.remove(bdr)
        st.paragraph_format.space_before = Pt(14 if kind != "title" else 0)
        st.paragraph_format.space_after = Pt(6)
    bullet_style = styles["List Bullet"]
    _set_style_font(bullet_style, body_font, body_size, color=body_color)

    style_for = {"title": "Title", "h1": "Heading 1", "h2": "Heading 2", "bullet": "List Bullet", "body": "Normal"}
    for p in paras:
        kind, ls = p["kind"], p["lines"]
        par = doc.add_paragraph(style=style_for[kind])
        if all(l.centered for l in ls):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if kind in ("title", "h1", "h2"):
            par.add_run(" ".join(l.text for l in ls))
            continue
        runs = []  # [text, bold, italic]
        for idx, l in enumerate(ls):
            for j, r in enumerate(l.runs()):
                text = r[0]
                if j == 0:
                    text = text.lstrip()
                    if idx > 0 and runs:
                        if runs[-1][0].endswith("-") and text[:1].islower():
                            runs[-1][0] = runs[-1][0][:-1]      # re-join a hyphenated word
                        elif not runs[-1][0].endswith(" "):
                            text = " " + text                     # line break becomes a space
                if runs and runs[-1][1] == r[1] and runs[-1][2] == r[2]:
                    runs[-1][0] += text
                else:
                    runs.append([text, r[1], r[2]])
        if kind == "bullet" and runs:
            # The bullet glyph is often its own run (e.g. in a symbol font); drop it, Word adds its own.
            runs[0][0] = BULLET_STRIP.sub("", runs[0][0].lstrip())
            while runs and not runs[0][0].strip():
                runs.pop(0)
            if runs:
                runs[0][0] = runs[0][0].lstrip()
        first = ls[0]
        for text, b, i in runs:
            if not text:
                continue
            run = par.add_run(text)
            if b:
                run.bold = True
            if i:
                run.italic = True
            if kind == "body":
                # Keep the look of one-off lines (cover subtitles, callouts) that differ from body text.
                if abs(first.size - body_size) >= 1.0:
                    run.font.size = Pt(first.size)
                if first.font != body_font:
                    run.font.name = first.font
                if first.color and first.color != body_color:
                    run.font.color.rgb = RGBColor.from_string(first.color)

    buf = io.BytesIO()
    doc.save(buf)
    notes = ["PDF converted to an editable Word layout. Images, tables, columns and exact positioning "
             "are not carried over; for the most faithful result use the .docx version of your sample."]
    return buf.getvalue(), notes
